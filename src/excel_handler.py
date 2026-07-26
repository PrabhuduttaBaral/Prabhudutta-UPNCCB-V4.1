"""
Prabhudutta-UPNCCB V4.1
excel_handler.py
"""

from openpyxl import load_workbook
from logger import logger


class ExcelHandler:

    def __init__(self, file_path):
        self.file_path = file_path
        self.workbook = None
        self.sheet = None

    def open(self):
        logger.info("Opening Excel File...")
        self.workbook = load_workbook(self.file_path)
        self.sheet = self.workbook.active
        logger.info("Excel File Opened Successfully")

    def total_records(self):
        return self.sheet.max_row - 1

    def get_record(self, row):
        """
        Column Mapping

        A = Account Number
        B = Amount
        C = Token Number

        Change these later if your Excel format is different.
        """

        account = self.sheet[f"A{row}"].value
        amount = self.sheet[f"B{row}"].value
        token = self.sheet[f"C{row}"].value

        return account, amount, token

    def update_voucher(self, row, voucher):
        """
        Writes Voucher Number in Column D
        """

        self.sheet[f"D{row}"] = voucher

    def save(self):
        self.workbook.save(self.file_path)
        logger.info("Excel Saved Successfully")

    def close(self):
        if self.workbook:
            self.workbook.close()